"""PowerPoint extraction: one block per slide, including tables, speaker notes and
the text inside pictures (screenshots and diagrams are very common in decks)."""
from __future__ import annotations

from pathlib import Path

from app.ingestion import ocr

from .common import clean, table_to_text

PICTURE_SHAPE = 13          # MSO_SHAPE_TYPE.PICTURE


def extract_pptx(path: Path) -> list[dict]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    budget = ocr.ImageBudget() if ocr.available() else None
    blocks: list[dict] = []

    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = clean(slide.shapes.title.text)
        except Exception:                                     # noqa: BLE001
            pass

        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape, title, budget))

        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = clean(slide.notes_slide.notes_text_frame.text)
                if notes:
                    parts.append(f"Speaker notes: {notes}")
        except Exception:                                     # noqa: BLE001
            pass

        text = clean("\n".join(parts))
        if title or text:
            blocks.append({"text": clean(f"{title}\n{text}"),
                           "loc": f"slide {number}", "heading": title})
    return blocks


def _shape_text(shape, title: str, budget: ocr.ImageBudget | None) -> list[str]:
    """Text of one shape: table, text frame, grouped shapes, or OCR of a picture."""
    out: list[str] = []
    try:
        if shape.shape_type == 6 and hasattr(shape, "shapes"):        # grouped shapes
            for child in shape.shapes:
                out.extend(_shape_text(child, title, budget))
            return out

        if getattr(shape, "has_table", False):
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            text = table_to_text(rows)
            if text:
                out.append(text)
            return out

        if getattr(shape, "has_text_frame", False):
            text = clean(shape.text)
            if text and text != title:
                out.append(text)

        if shape.shape_type == PICTURE_SHAPE and budget is not None:
            try:
                text = budget.read(shape.image.blob)
            except Exception:                                 # noqa: BLE001
                text = ""
            if text:
                out.append(f"[image] {text}")
    except Exception:                                         # noqa: BLE001
        return out
    return out


__all__ = ["extract_pptx"]
