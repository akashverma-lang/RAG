"""Excel and delimited-text extraction.

Rows are emitted as "column: value" records in batches, so a chunk always carries
its own headers - a retrieved row still makes sense without the rest of the sheet.
"""
from __future__ import annotations

import csv
from pathlib import Path

from .common import ROWS_PER_BLOCK, ocr_media_blocks, rows_block


def extract_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[dict] = []
    try:
        for sheet in workbook.worksheets:
            header: list[str] = []
            batch: list[list[str]] = []
            start = 2
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if v is None else str(v).strip() for v in row]
                if not any(values):
                    continue
                if not header:
                    header, start = values, index + 1
                    continue
                batch.append(values)
                if len(batch) >= ROWS_PER_BLOCK:
                    block = rows_block(batch, sheet.title, start, header)
                    if block["text"]:
                        blocks.append(block)
                    start += len(batch)
                    batch = []
            if batch:
                block = rows_block(batch, sheet.title, start, header)
                if block["text"]:
                    blocks.append(block)
            if header and not blocks:
                blocks.append({"text": " | ".join(header), "loc": sheet.title,
                               "heading": sheet.title})
    finally:
        workbook.close()

    blocks.extend(ocr_media_blocks(path))       # charts pasted in as pictures
    return blocks


def extract_xls(path: Path) -> list[dict]:
    try:
        import xlrd
    except ImportError:
        from .legacy import via_office_com

        return via_office_com(path, ".xlsx")

    book = xlrd.open_workbook(str(path))
    blocks: list[dict] = []
    for sheet in book.sheets():
        header: list[str] = []
        batch: list[list[str]] = []
        start = 2
        for r in range(sheet.nrows):
            values = ["" if v is None else str(v).strip() for v in sheet.row_values(r)]
            if not any(values):
                continue
            if not header:
                header, start = values, r + 2
                continue
            batch.append(values)
            if len(batch) >= ROWS_PER_BLOCK:
                blocks.append(rows_block(batch, sheet.name, start, header))
                start += len(batch)
                batch = []
        if batch:
            blocks.append(rows_block(batch, sheet.name, start, header))
    return [b for b in blocks if b["text"]]


def extract_csv(path: Path) -> list[dict]:
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        sample = fh.read(64 * 1024)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except Exception:                                     # noqa: BLE001
            dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel

        reader = csv.reader(fh, dialect)
        blocks: list[dict] = []
        header: list[str] = []
        batch: list[list[str]] = []
        start = 2
        for index, row in enumerate(reader, start=1):
            values = [c.strip() for c in row]
            if not any(values):
                continue
            if not header:
                header, start = values, index + 1
                continue
            batch.append(values)
            if len(batch) >= ROWS_PER_BLOCK:
                blocks.append(rows_block(batch, path.stem, start, header))
                start += len(batch)
                batch = []
        if batch:
            blocks.append(rows_block(batch, path.stem, start, header))
        if header and not blocks:
            blocks.append({"text": " | ".join(header), "loc": "header",
                           "heading": path.stem})
    return [b for b in blocks if b["text"]]


__all__ = ["extract_xlsx", "extract_xls", "extract_csv"]
