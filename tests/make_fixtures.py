"""Generate a small sample document set covering every supported format.

Used by ``python rag.py test``.  Some of these files carry text *only inside
pictures*, which is what exercises the OCR path.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path


# --------------------------------------------------------------------------- helpers
def render_png(lines: list[str], width: int = 620, height: int = 300,
               size: float = 26) -> bytes:
    """Render text to a clean PNG using real fonts, so OCR has a fair chance."""
    import pymupdf as fitz

    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    y = 60
    for line in lines:
        page.insert_text((40, y), line, fontsize=size, fontname="helv")
        y += size + 18
    pix = page.get_pixmap(dpi=150)
    data = pix.tobytes("png")
    doc.close()
    return data


# --------------------------------------------------------------------------- builder
def build(out: Path) -> Path:
    out = Path(out)
    (out / "reports").mkdir(parents=True, exist_ok=True)
    (out / "sheets").mkdir(parents=True, exist_ok=True)
    (out / "scans").mkdir(parents=True, exist_ok=True)

    import pymupdf as fitz

    # ---------------------------------------------------------------- pdf (text)
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((60, 90), "Quarterly Report 2025", fontsize=20)
    p1.insert_text((60, 140), "Acme Corp revenue for Q3 2025 was 4.7 million EUR,", fontsize=11)
    p1.insert_text((60, 160), "an increase of 18 percent over Q2 2025.", fontsize=11)
    p1.insert_text((60, 180), "The Berlin office contributed 1.2 million EUR.", fontsize=11)
    p2 = doc.new_page()
    p2.insert_text((60, 90), "Risks", fontsize=16)
    p2.insert_text((60, 130), "The main risk identified is supplier concentration:", fontsize=11)
    p2.insert_text((60, 150), "72 percent of components come from a single vendor,", fontsize=11)
    p2.insert_text((60, 170), "Nordwind GmbH. Mitigation deadline is 2026-03-31.", fontsize=11)
    doc.save(out / "reports" / "quarterly_report.pdf")
    doc.close()

    # ------------------------------------------------- pdf with a picture on a text page
    chart = render_png(["Q4 FORECAST", "Projected revenue 6.1 million EUR",
                        "Confidence level 82 percent"])
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((60, 80), "Forecast Appendix", fontsize=18)
    page.insert_text((60, 120), "The chart below summarises the Q4 outlook.", fontsize=11)
    page.insert_image(fitz.Rect(60, 150, 520, 380), stream=chart)
    doc.save(out / "reports" / "forecast_appendix.pdf")
    doc.close()

    # ---------------------------------------------------------- scanned pdf (no text layer)
    scan = render_png(["SERVICE AGREEMENT", "Termination notice is 90 days",
                       "Contract reference SA-4417"], height=340)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(0, 0, 595, 400), stream=scan)
    doc.save(out / "scans" / "service_agreement_scan.pdf")
    doc.close()

    # ---------------------------------------------------------------- standalone image
    (out / "scans" / "invoice_photo.png").write_bytes(
        render_png(["INVOICE 88421", "Total due 12450 EUR",
                    "Vendor Nordwind GmbH", "Payment date 2026-02-28"], height=360))

    # ---------------------------------------------------------------- word
    import docx

    d = docx.Document()
    d.add_heading("Employee Handbook", 0)
    d.add_heading("Leave Policy", level=1)
    d.add_paragraph("Every full-time employee receives 28 days of paid annual leave per year.")
    d.add_paragraph("Requests must be submitted at least 14 days in advance to the line manager.")
    d.add_heading("Remote Work", level=1)
    d.add_paragraph("Employees may work remotely up to 3 days per week, subject to team approval.")
    table = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["Region", "Headcount", "Manager"],
                             ["Berlin", "42", "Lena Fischer"],
                             ["Warsaw", "17", "Piotr Nowak"]]):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    d.add_heading("Building Access", level=1)
    d.add_paragraph("The access card reader is shown below.")
    d.add_picture(io.BytesIO(render_png(["SAFETY CODE ZULU 7719",
                                         "Report incidents within 24 hours"])))
    d.save(out / "reports" / "employee_handbook.docx")

    # ---------------------------------------------------------------- powerpoint
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Product Roadmap 2026"
    s1.placeholders[1].text = ("Q1: launch Atlas mobile app\n"
                              "Q2: SOC 2 certification\n"
                              "Q3: expand to the Nordic market")
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Budget"
    s2.placeholders[1].text = "Total engineering budget for 2026 is 2.4 million EUR."
    s2.notes_slide.notes_text_frame.text = "Remind the board that hiring is frozen until March."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Market Position"
    s3.shapes.add_picture(io.BytesIO(render_png(["MARKET SHARE 34 PERCENT",
                                                 "Ranked second in the region"])),
                          Inches(1), Inches(2), width=Inches(6))
    prs.save(out / "reports" / "roadmap.pptx")

    # ---------------------------------------------------------------- excel
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Region", "Units", "Revenue EUR"])
    for r in [("January", "Berlin", 120, 45000), ("February", "Berlin", 145, 52300),
              ("March", "Warsaw", 98, 31200), ("April", "Nordic", 210, 88900)]:
        ws.append(list(r))
    ws2 = wb.create_sheet("Costs")
    ws2.append(["Category", "Amount EUR"])
    for r in [("Cloud hosting", 18400), ("Travel", 5200), ("Licences", 9800)]:
        ws2.append(list(r))
    wb.save(out / "sheets" / "sales_2025.xlsx")

    # ---------------------------------------------------------------- csv
    with open(out / "sheets" / "customers.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["customer", "country", "plan", "mrr_eur", "renewal"])
        for r in [["Globex", "Germany", "Enterprise", 4200, "2026-01-15"],
                  ["Initech", "Poland", "Pro", 900, "2025-11-30"],
                  ["Umbrella", "Sweden", "Enterprise", 5100, "2026-06-01"]]:
            w.writerow(r)

    # ---------------------------------------------------------------- text formats
    (out / "notes.md").write_text(
        "# Meeting Notes\n\n## 2025-09-02 Standup\n\n"
        "- Ana will finish the migration script by Friday\n"
        "- The database upgrade is blocked on the vendor licence\n\n"
        "## Decisions\n\nWe agreed to postpone the Nordic launch to Q3 2026.\n",
        encoding="utf-8")
    (out / "readme.txt").write_text(
        "Support hotline: +49 30 1234567. Office hours are 09:00 to 17:00 CET.\n",
        encoding="utf-8")
    (out / "policy.html").write_text(
        "<html><head><style>p{color:red}</style></head><body>"
        "<h1>Security Policy</h1><p>All laptops must use full disk encryption.</p>"
        "<p>Passwords rotate every 90 days.</p></body></html>", encoding="utf-8")
    (out / "config.json").write_text(json.dumps(
        {"service": "atlas", "retention_days": 400, "owner": "platform-team"}, indent=2),
        encoding="utf-8")
    (out / "mail.eml").write_text(
        "From: lena@acme.example\nTo: board@acme.example\nSubject: Nordwind contract\n"
        "Date: Mon, 1 Sep 2025 10:00:00 +0200\nContent-Type: text/plain; charset=utf-8\n\n"
        "The Nordwind supplier contract expires on 2026-03-31 and must be renegotiated.\n",
        encoding="utf-8")
    return out


if __name__ == "__main__":
    import sys

    target = build(Path(sys.argv[1] if len(sys.argv) > 1 else "./sample_data"))
    print("fixtures written to", target)
    for f in sorted(target.rglob("*")):
        if f.is_file():
            print("  ", f.relative_to(target).as_posix(), f.stat().st_size, "bytes")
